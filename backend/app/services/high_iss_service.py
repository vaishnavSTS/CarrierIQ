"""STS's High ISS list: carriers whose latest Estimated ISS is at or above a threshold (75 by
default), worst first, with a phone number to call; also as an Excel file.

The scores come from iss_snapshots: once a month the worker rebuilds FMCSA's ISS tables and
recalculates every stored carrier (iss_service.snapshot_all), saving a snapshot when a score
changes. So the newest snapshot per carrier is its current score, and the list follows the
monthly recalculation on its own. Only carriers stored in CarrierSafetyScan are listed.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Address, Carrier, IssSnapshot, Phone
from app.models.enums import AddressType, PhoneType
from app.schemas.high_iss import HighIssCarrierOut, HighIssOut
from app.services.iss_service import latest_reference


def _pretty(digits: str) -> str:
    d = digits[-10:]
    return f"({d[:3]}) {d[3:6]}-{d[6:]}" if len(d) == 10 else digits


def high_iss(db: Session, min_score: int = 75) -> HighIssOut:
    latest = (
        select(IssSnapshot)
        .distinct(IssSnapshot.carrier_id)
        .order_by(IssSnapshot.carrier_id, IssSnapshot.created_at.desc(), IssSnapshot.id.desc())
        .subquery()
    )
    rows = db.execute(
        select(Carrier, latest.c.value, latest.c.band, latest.c.data_month)
        .join(latest, latest.c.carrier_id == Carrier.id)
        .where(latest.c.value >= min_score)
        .order_by(latest.c.value.desc(), Carrier.legal_name)
    ).all()
    carrier_ids = [carrier.id for carrier, *_ in rows]

    # One number per carrier: a current office phone first, then a current cell, then the newest.
    phones: dict[int, str] = {}
    for phone in db.scalars(
        select(Phone)
        .where(Phone.carrier_id.in_(carrier_ids), Phone.phone_type != PhoneType.FAX)
        .order_by(
            Phone.is_current.desc(),
            (Phone.phone_type == PhoneType.OFFICE).desc(),
            Phone.last_seen.desc(),
        )
    ):
        phones.setdefault(phone.carrier_id, phone.number_normalized)
    places: dict[int, str] = {}
    for address in db.scalars(
        select(Address).where(
            Address.carrier_id.in_(carrier_ids),
            Address.address_type == AddressType.PHYSICAL,
            Address.is_current.is_(True),
        )
    ):
        places[address.carrier_id] = ", ".join(p for p in [address.city, address.state] if p)

    reference = latest_reference(db)
    return HighIssOut(
        min_score=min_score,
        data_month=reference.data_month if reference else None,
        carriers=[
            HighIssCarrierOut(
                usdot_number=carrier.usdot_number,
                legal_name=carrier.legal_name,
                dba_name=carrier.dba_name,
                phone=_pretty(phones[carrier.id]) if carrier.id in phones else None,
                location=places.get(carrier.id) or None,
                fleet_size=carrier.fleet_size,
                iss=value,
                band=band,
                score_since=score_month,
            )
            for carrier, value, band, score_month in rows
        ],
    )


COLUMNS = [
    ("ISS score", 10),
    ("Carrier name", 40),
    ("DBA", 28),
    ("USDOT", 12),
    ("Phone", 16),
    ("Location", 24),
    ("Trucks", 9),
    ("ISS band", 11),
    ("Score since", 13),
]


def high_iss_xlsx(result: HighIssOut) -> bytes:
    """The list as an Excel workbook: one sheet, a bold frozen header, worst first."""
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    month = f"{result.data_month:%B %Y}" if result.data_month else "no ISS data yet"
    sheet.title = f"ISS {result.min_score}+"
    sheet.append([f"Carriers with Estimated ISS {result.min_score} or higher, FMCSA data {month}"])
    sheet["A1"].font = Font(bold=True, size=12)
    sheet.append([])
    sheet.append([name for name, _ in COLUMNS])
    for cell in sheet[3]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="EDE9FE")
    for c in result.carriers:
        sheet.append(
            [
                c.iss,
                c.legal_name,
                c.dba_name,
                c.usdot_number,
                c.phone,
                c.location,
                c.fleet_size,
                c.band,
                c.score_since,
            ]
        )
        sheet.cell(row=sheet.max_row, column=9).number_format = "mmm yyyy"
    for i, (_, width) in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    sheet.freeze_panes = "A4"
    out = BytesIO()
    book.save(out)
    return out.getvalue()
