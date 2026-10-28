"""STS's High ISS list: each carrier's latest Estimated ISS at or above the threshold, worst first,
with a phone to call; and the same list as an Excel file."""

from collections.abc import Iterator
from datetime import UTC, date, datetime, timedelta
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.db.session import get_db
from app.main import app
from app.models import Address, IssReference, IssSnapshot, Phone
from app.models.enums import AddressType, PhoneType
from tests.factories import make_carrier, make_raw_record

OCT = date(2026, 10, 1)


@pytest.fixture
def api(db: Session) -> Iterator[TestClient]:
    app.dependency_overrides[get_db] = lambda: db
    with TestClient(app) as test_client:
        yield test_client
    app.dependency_overrides.clear()


def snapshot(db: Session, carrier_id: int, value: int | None, month: date, age_days: int = 0):
    db.add(
        IssSnapshot(
            carrier_id=carrier_id,
            value=value,
            band="INSPECT" if value and value >= 75 else "OPTIONAL",
            group=1,
            algorithm="test",
            reason="test",
            data_month=month,
            version="iss-est test",
            basics=[],
            created_at=datetime.now(UTC) - timedelta(days=age_days),
        )
    )


@pytest.fixture
def carriers(db: Session) -> None:
    raw = make_raw_record(db)
    worst = make_carrier(db, 1000001, "WORST TRUCKING LLC")
    worst.fleet_size = 12
    high = make_carrier(db, 1000002, "HIGH HAULERS INC")
    improved = make_carrier(db, 1000003, "BETTER NOW LLC")
    low = make_carrier(db, 1000004, "LOW RISK LLC")
    snapshot(db, worst.id, 91, OCT)
    snapshot(db, high.id, 75, date(2026, 9, 1), age_days=40)  # unchanged since September
    snapshot(db, improved.id, 88, date(2026, 9, 1), age_days=40)
    snapshot(db, improved.id, 60, OCT)  # the latest counts
    snapshot(db, low.id, 40, OCT)
    seen = {"first_seen": date(2026, 1, 1), "last_seen": date(2026, 10, 1)}
    trace = {"source": raw.source, "raw_record_id": raw.id}
    db.add_all(
        [
            Phone(
                carrier_id=worst.id,
                phone_type=PhoneType.FAX,
                number_normalized="2095550000",
                **seen,
                **trace,
            ),
            Phone(
                carrier_id=worst.id,
                phone_type=PhoneType.OFFICE,
                number_normalized="2095550123",
                **seen,
                **trace,
            ),
            Address(
                carrier_id=worst.id,
                address_type=AddressType.PHYSICAL,
                city="TRACY",
                state="CA",
                **seen,
                **trace,
            ),
            IssReference(
                data_month=OCT, version="iss-est test", dataset_id="x", carriers=1, tables={}
            ),
        ]
    )
    db.flush()


def test_the_list_has_each_carriers_latest_score_75_and_up(api: TestClient, carriers: None) -> None:
    body = api.get("/api/v1/iss/high").json()
    assert body["min_score"] == 75 and body["data_month"] == "2026-10-01"
    assert [(c["legal_name"], c["iss"]) for c in body["carriers"]] == [
        ("WORST TRUCKING LLC", 91),
        ("HIGH HAULERS INC", 75),
    ]
    worst = body["carriers"][0]
    assert (worst["usdot_number"], worst["phone"], worst["location"], worst["fleet_size"]) == (
        1000001,
        "(209) 555-0123",  # the office phone, not the fax
        "TRACY, CA",
        12,
    )
    assert body["carriers"][1]["score_since"] == "2026-09-01"
    assert [c["iss"] for c in api.get("/api/v1/iss/high?min_score=90").json()["carriers"]] == [91]
    assert api.get("/api/v1/iss/high?min_score=0").status_code == 422


def test_the_excel_file(api: TestClient, carriers: None) -> None:
    answer = api.get("/api/v1/iss/high.xlsx")
    assert answer.status_code == 200
    assert answer.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "high-safety-score-75-plus-" in answer.headers["content-disposition"]
    sheet = load_workbook(BytesIO(answer.content)).active
    assert sheet is not None
    assert "October 2026" in str(sheet["A1"].value)
    assert [c.value for c in sheet[3]][:5] == ["ISS score", "Carrier name", "DBA", "USDOT", "Phone"]
    assert [(r[0], r[1], r[3]) for r in sheet.iter_rows(min_row=4, values_only=True)] == [
        (91, "WORST TRUCKING LLC", 1000001),
        (75, "HIGH HAULERS INC", 1000002),
    ]


def test_staff_only(api: TestClient, carriers: None, real_auth: None) -> None:
    assert api.get("/api/v1/iss/high").status_code == 401
    assert api.get("/api/v1/iss/high.xlsx").status_code == 401
